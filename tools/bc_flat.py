# -*- coding: utf-8 -*-
"""
bc_flat.py — '평탄화된' PPTX(표·그룹이 도형/텍스트박스로 풀린 파일) 복원기
  - flat_table(flat, prs)   : 도형 좌표로 표 격자 복원 → {'cols', 'rows'} (bc_convert.table_to_grid 와 동일 포맷)
  - flat_overview(flat, prs): 개요 슬라이드에서 (카테고리명, 상품 3줄) 쌍 복원
  - flat_reason(flat, prs)  : 선정이유 슬라이드에서 라벨(상품 특징/선정 이유/세일즈 포인트)별 내용 복원
좌표는 슬라이드 폭/높이 대비 %(0~100) 로 정규화해서 다룬다.
"""
import re
from pptx.enum.shapes import MSO_SHAPE_TYPE

LABELS = {'상품특징': 'feature', '선정이유': 'reason', '세일즈포인트': 'sales'}
DASH = r'[–\-—]'

def _norm(s): return re.sub(r'[\s/／·ㆍ,\.\n]+', '', s or '')
def _txt(s):
    if not getattr(s, 'has_text_frame', False) or not s.has_text_frame: return ''
    return '\n'.join(p.text.strip() for p in s.text_frame.paragraphs if p.text.strip())
def _paras(s):
    if not getattr(s, 'has_text_frame', False) or not s.has_text_frame: return []
    return [p.text.strip() for p in s.text_frame.paragraphs if p.text.strip()]

def _box(s, W, H):
    return (s.left / W * 100, s.top / H * 100, (s.left + s.width) / W * 100, (s.top + s.height) / H * 100)

def _cluster(vals, tol):
    vals = sorted(vals); out = []
    for v in vals:
        if out and v - out[-1][-1] <= tol: out[-1].append(v)
        else: out.append([v])
    return [sum(c) / len(c) for c in out]

def _near(bounds, v, tol):
    best = min(range(len(bounds)), key=lambda i: abs(bounds[i] - v))
    return best if abs(bounds[best] - v) <= tol else None

# ─────────────────────────────────────────────────────────────────────────────
def flat_table(flat, prs, y_min=12.0, y_max=91.0, tol=0.9):
    """도형 좌표로 표 복원. 표 영역(컨테이너) 안의 사각형/텍스트박스 모서리로 행·열 경계를 만든다."""
    W, H = prs.slide_width, prs.slide_height
    items = []
    for s in flat:
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE: continue
        x0, y0, x1, y1 = _box(s, W, H)
        if y0 < y_min or y1 > y_max + 1: continue
        items.append((s, x0, y0, x1, y1, _txt(s)))
    if not items: return None
    # 컨테이너(가장 큰 빈 사각형) → 표 영역
    big = [it for it in items if (it[3] - it[1]) > 60 and (it[4] - it[2]) > 30 and not it[5]]
    if big:
        cont = max(big, key=lambda it: (it[3] - it[1]) * (it[4] - it[2]))
        rx0, ry0, rx1, ry1 = cont[1] - tol, cont[2] - tol, cont[3] + tol, cont[4] + tol
        items = [it for it in items if it is not cont and it[1] >= rx0 and it[2] >= ry0 and it[3] <= rx1 and it[4] <= ry1]
    else:
        # 컨테이너가 없으면 텍스트가 있는 셀 박스들의 외곽을 표 영역으로
        cells0 = [it for it in items if it[5] and (it[3] - it[1]) >= 2 and (it[4] - it[2]) >= 2]
        if len(cells0) < 6: return None
        rx0 = min(it[1] for it in cells0) - tol; rx1 = max(it[3] for it in cells0) + tol
        ry0 = min(it[2] for it in cells0) - tol; ry1 = max(it[4] for it in cells0) + tol
        items = [it for it in items if it[1] >= rx0 and it[2] >= ry0 and it[3] <= rx1 and it[4] <= ry1]
    # 셀 후보: 선(두께<1%)이 아닌 사각형/텍스트박스
    cells = [it for it in items if (it[3] - it[1]) >= 1.5 and (it[4] - it[2]) >= 1.5]
    if len(cells) < 6: return None
    # 경계: 셀 후보들의 좌/우, 상/하 모서리 군집 — 2개 이상 셀이 공유하는 값만 경계로 인정
    def bounds(vals):
        cl = _cluster(vals, tol)
        cnt = [sum(1 for v in vals if abs(v - c) <= tol) for c in cl]
        return [c for c, n in zip(cl, cnt) if n >= 2]
    # 경계 기준 도형: 배경 사각형(텍스트 없음)이 충분하면 그것만 사용 — 셀 안 보조 텍스트가 경계를 오염시키지 않도록
    bg = [it for it in cells if not it[5]]
    src = bg if len(bg) >= 6 else cells
    xb = bounds([it[1] for it in src] + [it[3] for it in src])
    yb = bounds([it[2] for it in src] + [it[4] for it in src])
    if len(xb) < 3 or len(yb) < 3: return None
    ncol, nrow = len(xb) - 1, len(yb) - 1
    grid = {}      # (r0,c0,r1,c1) -> [texts]
    floating = []  # 경계에 맞지 않는 작은 텍스트(셀 안의 보조 텍스트)
    for it in cells:
        s, x0, y0, x1, y1, t = it
        c0, c1, r0, r1 = _near(xb, x0, tol), _near(xb, x1, tol), _near(yb, y0, tol), _near(yb, y1, tol)
        if None in (c0, c1, r0, r1) or c1 <= c0 or r1 <= r0:
            if t: floating.append(it)
            continue
        key = (r0, c0, r1, c1)
        grid.setdefault(key, [])
        if t: grid[key].append((y0, x0, t))
    # 겹치는 범위 정리: 같은 (r0,c0) 에 여러 span 이 있으면 텍스트가 있는/큰 것을 우선
    by_start = {}
    for key, ts in grid.items():
        r0, c0, r1, c1 = key
        cur = by_start.get((r0, c0))
        score = (1 if ts else 0, (r1 - r0) * (c1 - c0))
        if cur is None or score > cur[0]: by_start[(r0, c0)] = (score, key, ts)
    # 셀 안의 보조 텍스트를 중심점 기준으로 소속 셀에 합치기
    def find_cell(cx, cy):
        for (r0, c0), (_, key, ts) in by_start.items():
            _, _, r1, c1 = key
            if xb[c0] - tol <= cx <= xb[c1] + tol and yb[r0] - tol <= cy <= yb[r1] + tol: return (r0, c0)
        return None
    for s, x0, y0, x1, y1, t in floating:
        k = find_cell((x0 + x1) / 2, (y0 + y1) / 2)
        if k: by_start[k][2].append((y0, x0, t))
    # 점유 맵 → 빈 칸 채우기 → rows 생성
    occ = [[False] * ncol for _ in range(nrow)]
    rows = [[] for _ in range(nrow)]
    for (r0, c0), (_, key, ts) in sorted(by_start.items()):
        _, _, r1, c1 = key
        if any(occ[r][c] for r in range(r0, r1) for c in range(c0, c1)):
            continue   # 이미 점유된 영역과 겹치면 무시(배경 사각형 중복 등)
        for r in range(r0, r1):
            for c in range(c0, c1): occ[r][c] = True
        text = '\n'.join(t for _, _, t in sorted(ts))
        d = {'t': text, 'c': c0}
        if r1 - r0 > 1: d['h'] = r1 - r0
        if c1 - c0 > 1: d['w'] = c1 - c0
        rows[r0].append(d)
    for r in range(nrow):
        for c in range(ncol):
            if not occ[r][c]:
                rows[r].append({'t': '', 'c': c}); occ[r][c] = True
        rows[r].sort(key=lambda d: d['c'])
    # 모든 셀이 비어 있는 행 제거(장식용 띠 등)
    keep = [i for i, r in enumerate(rows) if any(d['t'] for d in r)]
    if len(keep) < len(rows):
        rows = [rows[i] for i in keep]   # 행 삭제 시 rowspan 이 어긋날 수 있으나, 빈 행은 통상 표 바깥 장식
    return {'cols': ncol, 'rows': rows, 'bounds': (xb[0], yb[0], xb[-1], yb[-1])}

# ─────────────────────────────────────────────────────────────────────────────
def flat_overview(flat, prs, tol=1.0):
    """개요 슬라이드: 좁은 사각형(카테고리) + 같은 높이의 넓은 사각형(상품 3줄) 쌍을 찾는다."""
    W, H = prs.slide_width, prs.slide_height
    rects, texts = [], []
    for s in flat:
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE: continue
        x0, y0, x1, y1 = _box(s, W, H); t = _txt(s)
        if y0 < 11 or y0 > 92: continue
        if t: texts.append((x0, y0, x1, y1, t, s))
        if not t and (x1 - x0) >= 6 and (y1 - y0) >= 6: rects.append((x0, y0, x1, y1))
    narrow = [r for r in rects if 6 <= (r[2] - r[0]) <= 22 and (r[3] - r[1]) <= 25]
    out = []
    for n in sorted(narrow, key=lambda r: r[1]):
        wide = [r for r in rects if (r[2] - r[0]) > 35 and abs(r[1] - n[1]) <= tol and r[0] >= n[2] - tol]
        if not wide: continue
        w = min(wide, key=lambda r: r[0])
        def inside(tb, r):
            cx, cy = (tb[0] + tb[2]) / 2, (tb[1] + tb[3]) / 2
            return r[0] - tol <= cx <= r[2] + tol and r[1] - tol <= cy <= r[3] + tol
        cat = ' '.join(tb[4].replace('\n', ' ') for tb in sorted(texts, key=lambda tb: (tb[1], tb[0])) if inside(tb, n))
        prods = []
        for tb in sorted(texts, key=lambda tb: tb[1]):
            if inside(tb, w):
                for line in _paras(tb[5]):
                    if re.search(DASH, line): prods.append(re.sub(r'\s+', ' ', line).strip())
        cat = re.sub(r'\s*/\s*', '/', cat).strip()
        if cat and prods: out.append((cat, prods))
    return out

# ─────────────────────────────────────────────────────────────────────────────
def flat_reason(flat, prs):
    """선정이유 슬라이드: 왼쪽 라벨(상품 특징/선정 이유/세일즈 포인트) ↔ 오른쪽 내용 박스를 세로 위치로 매칭."""
    W, H = prs.slide_width, prs.slide_height
    lab, body = [], []
    for s in flat:
        t = _txt(s)
        if not t: continue
        x0, y0, x1, y1 = _box(s, W, H)
        if y0 < 11 or y0 > 92: continue
        if x0 < 20: lab.append([y0, y1, t])
        else: body.append((y0, y1, s))
    # 라벨 조각 합치기(세일즈/포인트 처럼 두 박스로 나뉜 경우)
    lab.sort(); merged = []
    for y0, y1, t in lab:
        # 바로 위 조각이 아직 완전한 라벨이 아니고(세일즈 → 세일즈포인트), 합친 결과가 라벨 앞부분과 맞을 때만 병합
        if merged and y0 - merged[-1][1] <= 2.5 and _norm(merged[-1][2]) not in LABELS \
           and any(k.startswith(_norm(merged[-1][2] + t)) for k in LABELS):
            merged[-1][1] = y1; merged[-1][2] += t
        else: merged.append([y0, y1, t])
    labels = [((y0 + y1) / 2, LABELS[_norm(t)]) for y0, y1, t in merged if _norm(t) in LABELS]
    if not labels: return {}
    detail = {}
    for y0, y1, s in sorted(body, key=lambda b: b[0]):
        cy = (y0 + y1) / 2
        key = min(labels, key=lambda l: abs(l[0] - cy))[1]
        detail.setdefault(key, []).extend(re.sub(r'\s+', ' ', p) for p in _paras(s))
    return detail
